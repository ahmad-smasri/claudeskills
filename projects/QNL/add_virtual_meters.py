"""Add the virtual metering layer to a building's ontology.

Drives off the BUILDINGS registry below: --code picks the tier matrix (the
CLIENT's answer, per building), the sheet paths, and the physical-meter overlap
table. QNL and SSC are registered. Reads and writes .csv or .xlsx.


Regenerates the layer from the tier matrix against the *live* ontology rather
than patching the hand-built workbook, so the room identifiers agree by
construction and rooms added after that workbook was made (L1-145) are picked
up for free.

Emits, in order:

  1. early declarations  - the para: classes, the two thermal units and the
     entity:Metering system node, all of which are referenced by later rows
  2. one virtual meter per (meter class x spatial entity) the tier matrix
     selects, each carrying the Dar Cairo five-predicate block. A pair a
     physical meter also measures is reported, not skipped: the two answer
     different questions and the gap between them is worth seeing
  3. para:contributionFraction on every unit fed by an AHU, except units sitting
     in a shaft, riser or ceiling void. It is a container for the unit's chilled
     water consumption, which QNL has no point for: the backend replaces the
     ContributionFraction series with its own calculation

Timeseries references for the meter points ARE written, one per point. Both
halves of the key are derivable without the calculation engine's register: the
tsid is Dar Cairo's token for the meter class, and the entityId is the space the
meter meters. They are still derived rather than confirmed, so the pending file
remains as the checklist to hand that team. contributionFraction takes the same
shape, its tsid fixed at "ContributionFraction" and its entityId read off the
unit's existing points.

    python3 projects/QNL/add_virtual_meters.py
"""

import argparse
import collections
import csv
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
# One derivation of the telemetry entity id, shared with the air terminal layer,
# so the two cannot drift. `entity_id` is already taken in this module for the
# map of ids read off a unit's existing points, hence the alias.
from add_air_terminal_points import entity_id as space_key
ONTOLOGY = ROOT / "projects/QNL/QNL_Ontology.csv"
WIDTH = 27

SITE = "entity:QF"
METERING = "entity:Metering"

# A room whose identifier matches this holds plant or cabling, not people. The
# rule is Dar Cairo's own reading of what a space meter is for; QNL matches
# nothing here today (its shaft-dwelling assets are FCUs, which no AHU feeds)
# but the rule travels to the next building.
SHAFT = re.compile(r"(Shaft|Riser|Ceiling-Void|Raised-Floor)", re.I)

# --- the tier matrix, from Sheet1 of the hand-built workbook -----------------
# B building, F floor, R room. Segment names are Dar Cairo's verbatim, except
# HW-Power-Thermal-Virtual-Meter, coined here to mirror the CHW one.
ELEC, THERMAL = "elec", "thermal"
METER_TYPES = [
    # class,                    segment,                               tiers,   kind
    ("para:Utility_Meter",      "Utility-Virtual-Meter",               "B",     ELEC),
    # para:UPS_Meter is deliberately absent from both matrices: a UPS meter
    # sums UPS load, and neither building publishes a UPS datapoint. Removed
    # from both sheets on 2026-09-09 at the client's direction. Do not put it
    # back without an input to sum - it would render an empty tile.
    # para:HW_Meter is deliberately absent: it measures DOMESTIC hot water and
    # neither building has an energy point for any. QNL's only DHW asset is a
    # calorifier carrying two alarms and a temperature, outside the selected
    # scope; SSC's heating is electric. Removed 2026-09-09, client direction.
    ("para:SPWR_Meter",         "SPWR-Util-Electrical-Virtual-Meter",  "BF",    ELEC),
    ("para:Common_Util_Meter",  "Common-Util-Electrical-Virtual-Meter","BF",    ELEC),
    ("para:CHW_Meter",          "CHW-Power-Thermal-Virtual-Meter",     "BFR",   THERMAL),
    ("para:HVAC_Meter",         "HVAC-Util-Electrical-Virtual-Meter",  "BFR",   ELEC),
    ("para:LTG_Meter",          "LTG-Util-Electrical-Virtual-Meter",   "BFR",   ELEC),
    ("brick:Electrical_Meter",  "Electrical-Virtual-Meter",            "BFR",   ELEC),
]

# Consumption is energy, Demand is power. Thermal carries para:KiloWt so a
# building rollup cannot silently add chilled-water kW to electrical kW.
POINTS = {
    ELEC: [
        ("Consumption", "brick:Electrical_Energy_Usage_Sensor", "unit:KiloW-HR"),
        ("Demand",      "brick:Electric_Power_Sensor",          "unit:KiloW"),
    ],
    THERMAL: [
        ("Consumption", "brick:Thermal_Energy_Usage_Sensor", "para:KiloWt-HR"),
        ("Demand",      "brick:Thermal_Power_Sensor",        "para:KiloWt"),
    ],
}

# Dar Cairo's timeseries token per (meter class, point kind). These are WRITTEN
# into the sheet, on the point's own ref:hasExternalReference row - Dar Cairo does
# the same (entity:Dar-Cairo_UPS-Util-Electrical-Virtual-Meter-Consumption carries
# UPS_KWH_CALC with para:hasEntityId "Smart Village"). The entityId half is the
# space the meter meters, underscored: QNL, QNL_L1, QNL_B_001A_Break_Out_Area.
# Dar Cairo puts one site constant on all of them, which only works where there
# is a single meter per class; QNL has 360 room-tier CHW meters that would
# collide on it. The pending file survives as the calculation-engine team's
# confirmation checklist, now with both halves filled rather than the entityId
# blank.
TSID = {
    ("para:Utility_Meter",     "Consumption"): "Utility_KWH",
    ("para:Utility_Meter",     "Demand"):      "Utility_KW",
    ("para:UPS_Meter",         "Consumption"): "UPS_KWH_CALC",
    ("para:UPS_Meter",         "Demand"):      "UPS_KW_CALC",
    ("para:SPWR_Meter",        "Consumption"): "SPWR_KWH_CALC",
    ("para:SPWR_Meter",        "Demand"):      "SPWR_KW_CALC",
    ("para:Common_Util_Meter", "Consumption"): "COMMON_KWH_CALC",
    ("para:Common_Util_Meter", "Demand"):      "COMMON_KW_CALC",
    ("para:HVAC_Meter",        "Consumption"): "HVAC_KWH_CALC",
    ("para:HVAC_Meter",        "Demand"):      "HVAC_KW_CALC",
    ("para:LTG_Meter",         "Consumption"): "LTG_KWH_CALC",
    ("para:LTG_Meter",         "Demand"):      "LTG_KW_CALC",
    ("brick:Electrical_Meter", "Consumption"): "ELEC_KWH_CALC",
    ("brick:Electrical_Meter", "Demand"):      "ELEC_KW_CALC",
    ("para:CHW_Meter",         "Consumption"): "CWPWR_KWHT_CALC",
    ("para:CHW_Meter",         "Demand"):      "CWPWR_KWT_CALC",
    ("para:HW_Meter",          "Consumption"): "HWPWR_KWHT_CALC",
    ("para:HW_Meter",          "Demand"):      "HWPWR_KWT_CALC",
}

# Physical meters do NOT suppress virtual ones. A physical meter reads what is
# actually imported at one point; a virtual meter is a calculated roll-up over a
# space. They answer different questions and can legitimately disagree - the gap
# between them is losses and unmetered load, which is worth seeing. So overlaps
# are REPORTED, never skipped, and the reviewer decides.
# (Client direction 2026-09-03, reversing an earlier rule that suppressed them.)
PHYSICAL_OVERLAP = {
    # (meter class, metered entity): the physical meter that also measures it
    ("para:Utility_Meter", "entity:QNL"): "entity:QNL_Total-Energy",
    ("para:CHW_Meter", "entity:QNL"): "entity:QNL_CHWS-MAIN-LOOP_Energy-Meter",
}

# --- per-building data ------------------------------------------------------
# The tier matrix is the CLIENT'S answer, not a house default, so it lives here
# per building rather than as one module constant. Same for the paths and the
# physical-meter overlaps, which are read off each building's own meter list.
#
# SSC (2026-09-09). Audited against SSC_Historian_IO_list_CP2.xlsx, 5,751 tags:
#   Utility    - SSC_MV_InACB, SSC_ELEC_MFM_MV_OG_I3.kW/.kWh, SSC_EnergyConsumptionCalc.kW
#   SPWR       - SPWR is SMALL POWER, not solar; Dar Cairo labels it "Small Power
#                Electrical Meter". 200 SMDB power/energy tags are the candidate
#                inputs, pending the electrical schedule to say which circuits.
#   Common     - same SMDB boards, same pending schedule
#   CHW        - SSC_CHWConsumption.KWh + 20 thermal points
#   HVAC       - 22 AHU kW/kWh, 8 CHW pump, 49 MCC power tags
#   Electrical - 100 SMDB kW, 100 SMDB MWh, 110 MV tags
#   LTG        - NO energy input: 55 SSC_LCPB_* circuits are all On/Off status.
#                Built at the client's direction to match QNL, which is in the
#                same position. Every one renders empty until a kWh tag exists.
#   UPS        - ELIMINATED: no inputs at all, no UPS datapoint in either
#                building. Client direction 2026-09-09.
#   HW         - ELIMINATED: SSC's heating is electric (5 AHU heater commands,
#                14 CRAC heater statuses). No hot-water loop, 0 hot-water tags.
SSC_MATRIX = [
    ("para:Utility_Meter",      "Utility-Virtual-Meter",               "B",     ELEC),
    ("para:SPWR_Meter",         "SPWR-Util-Electrical-Virtual-Meter",  "BF",    ELEC),
    ("para:Common_Util_Meter",  "Common-Util-Electrical-Virtual-Meter","BF",    ELEC),
    ("para:CHW_Meter",          "CHW-Power-Thermal-Virtual-Meter",     "BFR",   THERMAL),
    ("para:HVAC_Meter",         "HVAC-Util-Electrical-Virtual-Meter",  "BFR",   ELEC),
    ("para:LTG_Meter",          "LTG-Util-Electrical-Virtual-Meter",   "BFR",   ELEC),
    ("brick:Electrical_Meter",  "Electrical-Virtual-Meter",            "BFR",   ELEC),
]

BUILDINGS = {
    "QNL": {
        "ontology": ROOT / "projects/QNL/QNL_Ontology.csv",
        "out":      ROOT / "projects/QNL/QNL_Ontology.csv",
        "pending":  ROOT / "projects/QNL/QNL_virtual_meter_timeseries_pending.csv",
        "matrix":   None,          # filled with METER_TYPES below
        "overlap":  None,          # filled with PHYSICAL_OVERLAP below
        "contribution": True,      # client asked for it (QNL-036)
    },
    "SSC": {
        "ontology": ROOT / "reference-models/QF_SSC_Ontology_V04.xlsx",
        "out":      ROOT / "reference-models/QF_SSC_Ontology_V04.xlsx",
        "pending":  ROOT / "projects/SSC/SSC_virtual_meter_timeseries_pending.csv",
        "matrix":   SSC_MATRIX,
        # SSC's 45 existing meters are equipment-tier and NOT ONE carries a
        # brick:meters row, so the graph cannot say what any of them covers and
        # no overlap can be computed. Reported in the handover instead of guessed.
        "overlap":  {},
        # para:contributionFraction is a SEPARATE question the client has to be
        # asked (virtual-meters.md, "Ask first" question 2), not a side effect of
        # asking for meters. Asked and answered yes on 2026-09-09: it is what
        # gives SSC's 61 VAV-served rooms a real chilled-water input, and without
        # it the room-tier CHW meters there sum nothing. Added after the metering
        # layer had already landed, so it went in with --only-contribution.
        "contribution": True,
    },
}

DECLARATIONS = [
    ("para:Metering_System",     "brick:System",              "Metering System"),
    # QNL already carried this one, which is why it was originally omitted;
    # build_declarations() skips a class the sheet already declares, so naming
    # it here is a no-op for QNL and the difference between a working and a
    # dangling reference for any building that does not.
    ("para:Utility_Meter",       "brick:Electrical_Meter",    "Utility Electrical Meter"),
    ("para:SPWR_Meter",          "brick:Electrical_Meter",    "Small Power Electrical Meter"),
    ("para:Common_Util_Meter",   "brick:Electrical_Meter",    "Common Utilities Electrical Meter"),
    ("para:HVAC_Meter",          "brick:Electrical_Meter",    "HVAC Electrical Meter"),
    ("para:LTG_Meter",           "brick:Electrical_Meter",    "Lighting Electrical Meter"),
    ("para:CHW_Meter",           "brick:Thermal_Power_Meter", "Chilled Water Thermal Power Meter"),
    ("para:contributionFraction", "brick:Point",              "Contribution Fraction"),
]
UNIT_DECLARATIONS = [("para:KiloWt", "kWt"), ("para:KiloWt-HR", "kWt·hr")]

BUILDINGS["QNL"]["matrix"] = METER_TYPES
BUILDINGS["QNL"]["overlap"] = PHYSICAL_OVERLAP


# Which side a property belongs on is a property of the ROW, not of the property
# name: rdfs:label_en labels the meter on `meter isPartOf Metering` and labels the
# point on `meter hasPoint point`. Getting this backwards is the commonest
# authoring mistake here, so each caller names the side rather than sharing a
# lookup that cannot tell the two rows apart.
SUBJ_SLOTS = [5, 9, 13, 17, 21, 25]
OBJ_SLOTS = [7, 11, 15, 19, 23]


def row(subject, stype, predicate, obj="", otype="", sprops=(), oprops=()):
    """One sheet row. sprops describe the subject, oprops the object; each is a
    sequence of (column_name, value) pairs laid into the prop groups in order."""
    r = [subject, stype, predicate, obj, otype] + [""] * (WIDTH - 5)
    for slots, props in ((SUBJ_SLOTS, sprops), (OBJ_SLOTS, oprops)):
        for i, (name, value) in zip(slots, props):
            r[i], r[i + 1] = name, value
    return r


def label_of(identifier):
    """Label from the identifier, mechanically - one map so the two cannot drift."""
    return identifier.split(":", 1)[1].replace("_", " ").replace("-", " ")


def ontology_sheet(wb):
    """The triples sheet, picked by its header - never by tab name or .active."""
    head = ("subject", "subjecttype", "predicate", "object", "objecttype")
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        if tuple(str(c or "").strip().lower() for c in first[:5]) == head:
            return ws
    raise SystemExit("no ontology sheet in the workbook")


def read_ontology(path):
    """Rows from a .csv or .xlsx sheet, padded to the sheet width."""
    if str(path).lower().endswith(".xlsx"):
        import openpyxl
        ws = ontology_sheet(openpyxl.load_workbook(path, read_only=True, data_only=True))
        rows = [["" if c is None else str(c).strip() for c in r]
                for r in ws.iter_rows(values_only=True)]
        width = max(len(r) for r in rows)
        return [r + [""] * (width - len(r)) for r in rows]
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.reader(fh))


def write_ontology(path, header, body):
    """Write back in the format the sheet arrived in."""
    if str(path).lower().endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(header)
        for r in body:
            ws.append(r + [""] * (len(header) - len(r)) if len(r) < len(header) else r)
        wb.save(path)
        return
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(body)


def index(rows):
    """Everything the generator needs to know about the existing sheet."""
    etype, located, fedby, owner = {}, {}, collections.defaultdict(list), {}
    entity_id = collections.defaultdict(set)
    declared = set()
    for r in rows[1:]:
        if len(r) < 5:
            continue
        if r[0] and r[1]:
            etype.setdefault(r[0], r[1])
        if r[3] and r[4]:
            etype.setdefault(r[3], r[4])
        if r[1] == "owl:Class":
            declared.add(r[0])
        if r[2] == "rec:locatedIn":
            located[r[0]] = r[3]
        if r[2] == "rec:isFedBy":
            fedby[r[0]].append(r[3])
        if r[2] == "brick:hasPoint":
            owner[r[3]] = r[0]
    for r in rows[1:]:
        for i in range(5, len(r) - 1, 2):
            if r[i] == "para:hasEntityId" and r[i + 1] and r[0] in owner:
                entity_id[owner[r[0]]].add(r[i + 1])
    return etype, located, fedby, entity_id, declared


def spatial_targets(etype):
    """The building, its levels and its rooms, each with the type to write in
    the objectType column."""
    tiers = {"B": [], "F": [], "R": []}
    for e, t in etype.items():
        if t == "rec:Building":
            tiers["B"].append((e, t))
        elif t in ("rec:Level", "rec:BasementLevel", "rec:RoofLevel", "rec:GroundLevel"):
            tiers["F"].append((e, t))
        elif t == "rec:Room":
            tiers["R"].append((e, t))
    return {k: sorted(v) for k, v in tiers.items()}


def build_meters(tiers, matrix, overlap):
    """The metering layer. Returns (rows, pending timeseries, physical overlaps)."""
    out, pending, overlaps = [], [], []
    for cls, segment, applies, kind in matrix:
        for tier in applies:
            for target, target_type in tiers[tier]:
                covered = overlap.get((cls, target))
                if covered:
                    overlaps.append((f"{target}_{segment}", cls, covered))
                meter = f"{target}_{segment}"
                mlabel = label_of(meter)
                out.append(row(meter, cls, "brick:isPartOf", METERING, "para:Metering_System",
                               sprops=[("rdfs:label_en", mlabel)]))
                out.append(row(meter, cls, "brick:meters", target, target_type))
                out.append(row(meter, cls, "brick:isVirtualMeter", "<blanknode>", "<blanknode>",
                               oprops=[("brick:value", "TRUE")]))
                out.append(row(meter, cls, "rec:locatedIn", target, target_type))
                for kindname, pcls, unit in POINTS[kind]:
                    point = f"{meter}-{kindname}"
                    key = TSID.get((cls, kindname), "")
                    eid = space_key(target)
                    out.append(row(meter, cls, "brick:hasPoint", point, pcls,
                                   oprops=[("rdfs:label_en", f"{mlabel} {kindname}"),
                                           ("brick:hasUnit", unit)]))
                    # The point's own reference row. A meter block is EIGHT rows,
                    # not six: without this the point is the object of one
                    # hasPoint row and the subject of nothing, so the front end
                    # draws a tile with no series behind it.
                    if key:
                        out.append(row(point, pcls, "ref:hasExternalReference",
                                       "<blanknode>", "ref:TimeseriesReference",
                                       oprops=[("ref:hasTimeseriesId", key),
                                               ("para:hasEntityId", eid)]))
                    pending.append((point, pcls, key, eid, target))
    return out, pending, overlaps


def build_contribution(etype, located, fedby, entity_id):
    """para:contributionFraction on every AHU-fed unit outside a shaft."""
    out, skipped, derived = [], [], []
    fed = sorted(u for u, sources in fedby.items()
                 if any(etype.get(s) == "brick:Air_Handling_Unit" for s in sources))
    for unit in fed:
        room = located.get(unit, "")
        if room and SHAFT.search(room):
            skipped.append((unit, room))
            continue
        ids = entity_id.get(unit) or set()
        if len(ids) == 1:
            eid = next(iter(ids))
        else:
            # The rule is proven by every one of the 296 units that do carry a
            # key, so deriving the 297th is safe - but it is reported, not silent.
            eid = unit.split(":", 1)[1].replace("-", "_")
            derived.append((unit, eid))
        point = f"{unit}_Contribution-Fraction"
        out.append(row(unit, etype[unit], "brick:hasPoint", point, "para:contributionFraction",
                       oprops=[("rdfs:label_en", "Contribution Fraction"),
                               ("brick:hasUnit", "unit:UNITLESS")]))
        out.append(row(point, "para:contributionFraction", "ref:hasExternalReference",
                       "<blanknode>", "ref:TimeseriesReference",
                       oprops=[("ref:hasTimeseriesId", "ContributionFraction"),
                               ("para:hasEntityId", eid)]))
    return out, fed, skipped, derived


def build_declarations(declared, matrix, contribution):
    """Only what THIS building's matrix actually uses.

    Declaring everything in DECLARATIONS regardless of the matrix is how SSC
    ended up with a para:HW_Meter owl:Class row and no HW meter under it - a
    dangling class that reads as a modelling decision and is really a leak. The
    same bug had already put a para:UPS_Meter declaration in a sheet whose
    matrix never asked for one. The matrix is the authority.
    """
    needed = {"para:Metering_System"} | {cls for cls, _, _, _ in matrix}
    if contribution:
        needed.add("para:contributionFraction")
    out = []
    for cls, parent, label in DECLARATIONS:
        if cls in declared or cls not in needed:
            continue
        out.append(row(cls, "owl:Class", "rdfs:subClassOf", parent,
                       sprops=[("rdfs:label_en", label)]))
    # The thermal units follow the matrix too: a building with no THERMAL row
    # has no use for para:KiloWt and should not be handed one.
    if any(kind == THERMAL for *_, kind in matrix):
        for unit, symbol in UNIT_DECLARATIONS:
            out.append(row(unit, "qudt:Unit", "rdf:type", "qudt:Unit",
                           sprops=[("qudt:symbol", symbol)]))
    out.append(row(METERING, "para:Metering_System", "brick:isPartOf", SITE, "rec:Site",
                   sprops=[("rdfs:label_en", "Metering System")]))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--code", default="QNL", choices=sorted(BUILDINGS),
                    help="which building's tier matrix and paths to use")
    ap.add_argument("--ontology")
    ap.add_argument("--out")
    ap.add_argument("--pending")
    ap.add_argument("--only-contribution", action="store_true",
                    help="add para:contributionFraction alone, to a sheet whose\n"
                         "metering layer already landed. Writes no meters and does\n"
                         "not touch the pending file, which belongs to that layer.")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    b = BUILDINGS[args.code]
    ontology = args.ontology or str(b["ontology"])
    out_path = args.out or str(b["out"])
    pending_path = args.pending or str(b["pending"])

    rows = read_ontology(ontology)
    header = rows[0]
    etype, located, fedby, entity_id, declared = index(rows)

    # Narrowed to the SPACE-tier segments this script writes. The bare string
    # "Virtual-Meter" also matches equipment-tier meters a delivered sheet may
    # already carry - SSC has 27 of them, on AHU coil valves and FCUs - and those
    # are a different family that this layer neither replaces nor duplicates.
    segments = tuple(seg for _, seg, _, _ in b["matrix"])
    existing = {r[0] for r in rows[1:]
                if any(r[0].endswith("_" + g) for g in segments)}
    if existing and not args.only_contribution:
        raise SystemExit("sheet already carries a space-tier metering layer "
                         f"({len(existing)} meters) - rerun on a clean ontology")

    if args.only_contribution:
        # contributionFraction was decided separately from the metering layer on
        # both buildings, so it has to be addable to a sheet the layer already
        # landed on. The layer's own guard above is therefore skipped, and this
        # one takes its place: the points are what must not be written twice.
        if not b["contribution"]:
            raise SystemExit(f'{args.code} has contribution=False - set it before '
                             'asking for the points')
        already = {r[0] for r in rows[1:] if r[1] == "para:contributionFraction"}
        if already:
            raise SystemExit("sheet already carries para:contributionFraction on "
                             f"{len(already)} points - nothing to add")
        tiers = {"B": [], "F": [], "R": []}
        decls, meters, pending, overlaps = [], [], [], []
    else:
        tiers = spatial_targets(etype)
        decls = build_declarations(declared, b["matrix"], b["contribution"])
        meters, pending, overlaps = build_meters(tiers, b["matrix"], b["overlap"])

    if b["contribution"]:
        contrib, fed, skipped, derived = build_contribution(etype, located, fedby, entity_id)
    else:
        contrib, fed, skipped, derived = [], [], [], []

    print(f"spatial targets   building {len(tiers['B'])}  levels {len(tiers['F'])}  rooms {len(tiers['R'])}")
    print(f"declarations      {len(decls)} rows")
    print(f"virtual meters    {len(meters) // 8} meters, {len(meters)} rows")
    for m, cls, covered in overlaps:
        print(f"  note: {m} overlaps physical {covered} - both kept, they are not duplicates")
    if b["contribution"]:
        print(f"contributionFraction  {len(contrib) // 2} of {len(fed)} AHU-fed units"
              f"  (skipped in shafts: {len(skipped)})")
    else:
        print("contributionFraction  not requested for this building - skipped")
    if derived:
        print(f"  entityId derived rather than reused for: {[u for u, _ in derived]}")
    if skipped:
        for u, r in skipped:
            print(f"  skipped {u} in {r}")

    if args.dry_run:
        return

    body = rows[1:] + decls + meters + contrib
    # Declarations have to precede their first use, and the converter reads the
    # sheet top to bottom.
    body.sort(key=lambda r: 0 if r[1] in ("owl:Class", "qudt:Unit") else 1)
    write_ontology(out_path, header, body)
    print(f"wrote {out_path}  ({len(body)} rows, was {len(rows) - 1})")

    if args.only_contribution:
        # The pending file is the meter layer's checklist. A contribution-only
        # run has no meter keys to write and must not blank it.
        return

    with open(pending_path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["point", "point_class", "proposed_hasTimeseriesId",
                    "hasEntityId_TO_CONFIRM", "meters"])
        w.writerows(pending)
    print(f"wrote {pending_path}  ({len(pending)} derived keys for the calculation engine to confirm)")


if __name__ == "__main__":
    main()
