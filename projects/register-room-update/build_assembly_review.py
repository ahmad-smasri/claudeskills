#!/usr/bin/env python3
"""Write the RDC assembly review workbook.

One physical air-terminal unit is registered across several rows - the box, its
valve, and each air terminal it feeds - while the BMS screens draw it as a
single widget. This groups the rows back into assemblies and says, for each
one, whether its rows agree about the room.
"""
import collections
import csv
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import match_terminals as M

OUT = M.HERE / 'RDC_assembly_review.xlsx'
HEAD = PatternFill('solid', fgColor='FF1F4E79')
RED = PatternFill('solid', fgColor='FFFFC7CE')     # rooms differ
AMBER = PatternFill('solid', fgColor='FFFFE699')   # wording differs only
GREEN = PatternFill('solid', fgColor='FFC6EFCE')   # a screen reading applies


def head(ws, cols, row=1):
    for i, (name, w) in enumerate(cols, 1):
        c = ws.cell(row, i, name)
        c.fill, c.font = HEAD, Font(bold=True, color='FFFFFFFF')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 1)


def main():
    rows, read = M.load()
    partner = collections.defaultdict(set)
    for group in M.assemblies(rows):
        for t in group:
            partner[t] = group - {t}
    room = {t: d for _, t, _, d in rows}
    kind = {t: k for _, t, k, _ in rows}
    order = {t: i for i, t, _, _ in rows}

    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = 'START HERE'
    intro.column_dimensions['A'].width = 116
    text = [
        ('RDC - one unit, several register rows', True),
        ('', False),
        ('The RDC register lists one air-terminal assembly across several rows. '
         'A VAV box is `VAV4110_EV4111` - the box and its valve - and the air '
         'terminal it feeds is a row of its own, `AT-4110`. The BMS screens draw '
         'the whole assembly as one widget, labelled the way the register writes '
         'the box: `NB-VAV4110-EV4111`. No widget is ever labelled AT; there is '
         'nothing on the screen to click for a terminal.', False),
        ('', False),
        ('That is why 644 AT rows came back with no reading: they have no widget '
         'of their own, not because nothing serves them.', False),
        ('', False),
        ('How the rows are grouped', True),
        ('By the unit number, within one building, one level and one zone, and '
         'only across the air-terminal families (AT, EV, VEV, CEV, FEV, VAV, '
         'CAV, EAV). Three things that grouping has to get right:', False),
        ('  FCU and AHU are left out. FCU1004 and VAV1004 share a number and are '
         'different units - both are drawn as separate widgets on South GF-6.', False),
        ('  The south building registers terminals on L0 and boxes on GF. Those '
         'are one floor under two names, so they are treated as one; every other '
         'level is kept apart, or NB_1F_AT-5930 merges with NB_2F_AT-5930.', False),
        ('  VEV-5192a and VEV-5192b are two units. The letter is part of the '
         'number, not decoration.', False),
        ('', False),
        ('What the sheets say', True),
        ('  Assemblies    every register row, grouped. `rows in assembly` is how '
         'many rows the unit is split across.', False),
        ('  Conflicts     the assemblies whose rows do not agree about the room.', False),
        ('', False),
        ('Fills on the Assemblies sheet:', True),
        ('  red     the assembly names more than one room, and the room numbers '
         'differ - the register contradicts itself', False),
        ('  amber   same room, different wording (Conf against Conference, or a '
         'typo such as Director Ofice)', False),
        ('  green   a BMS screen reading applies to this row, either its own or '
         "its assembly's", False),
        ('', False),
        ('A box feeding several terminals is not a conflict. AT-7120 runs to nine '
         'of them, numbered -N1 to -N9, in nine different rooms, and those rows '
         "keep their own room rather than taking the box's.", False),
    ]
    for i, (line, bold) in enumerate(text, 1):
        c = intro.cell(i, 1, line)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if bold:
            c.font = Font(bold=True, size=12)

    ws = wb.create_sheet('Assemblies')
    head(ws, [('Assembly', 30), ('Rows in assembly', 9), ('Register tag', 30),
              ('Type', 8), ('Room in column D', 40), ('Assembly agrees?', 16),
              ('Room per BMS screen', 30), ('Read from', 26)])
    seen, r = set(), 1
    conflicts = []
    for _, tag, _, _ in rows:
        group = frozenset({tag} | partner[tag])
        if group in seen:
            continue
        seen.add(group)
        members = sorted(group, key=lambda t: order.get(t, 0))
        named = {t: room[t] for t in members if room.get(t)}
        many_terms = sum(1 for t in named if re.search(r'-N\d+$', t)) > 1
        keys = {M.norm(v) for v in named.values()}
        if many_terms or len(keys) < 2:
            state = 'yes'
        elif len({k[0] for k in keys}) > 1:
            state = 'different room'
        else:
            state = 'wording only'
        if state != 'yes':
            conflicts.append((state, members))
        label = re.sub(r'^RDC_', '', members[0])
        for t in members:
            r += 1
            hit = read.get(t)
            if not hit:
                sib = {p: read[p] for p in partner[t] if p in read}
                if len(sib) == 1 and not re.search(r'-N\d+$', t):
                    src = next(iter(sib))
                    hit = (sib[src][0], sib[src][1], False)
                    src_label = re.sub(r'^RDC_', '', src)
                else:
                    src_label = ''
            else:
                src_label = 'this row'
            ws.cell(r, 1, label if t == members[0] else '')
            ws.cell(r, 2, len(members) if t == members[0] else None)
            ws.cell(r, 3, t)
            ws.cell(r, 4, kind.get(t, ''))
            ws.cell(r, 5, room.get(t, ''))
            ws.cell(r, 6, state if t == members[0] else '')
            ws.cell(r, 7, hit[0] if hit else '')
            ws.cell(r, 8, src_label)
            fill = (RED if state == 'different room'
                    else AMBER if state == 'wording only'
                    else GREEN if hit else None)
            if fill:
                for c in range(1, 9):
                    ws.cell(r, c).fill = fill
    for rr in range(2, r + 1):
        for c in range(1, 9):
            ws.cell(rr, c).alignment = Alignment(vertical='top', wrap_text=True)

    cs = wb.create_sheet('Conflicts')
    head(cs, [('Kind', 16), ('Register tag', 32), ('Type', 8),
              ('Room in column D', 52)])
    cr = 1
    for state, members in sorted(conflicts, key=lambda c: c[0] != 'different room'):
        for t in members:
            cr += 1
            cs.cell(cr, 1, state)
            cs.cell(cr, 2, t)
            cs.cell(cr, 3, kind.get(t, ''))
            cs.cell(cr, 4, room.get(t, ''))
            if state == 'different room':
                for c in range(1, 5):
                    cs.cell(cr, c).fill = RED
        cr += 1
    wb.save(OUT)
    print('wrote %s - %d assemblies, %d with a conflict'
          % (OUT.name, len(seen), len(conflicts)))


if __name__ == '__main__':
    main()
