#!/usr/bin/env python3
"""Decide which RDC register rows are units and which are parts of one.

The test is the historian, not the tag family. A row is a controllable asset if
the historian carries an object for it - that is what "controllable" means, and
what the Controllable Asset Registry is for. A row with no historian object has
no points, cannot be commanded and cannot be read, and belongs to the ontology
as a `brick:hasPart` of the unit it is part of rather than as an entity of its
own.

Doing this by family instead would be wrong in both directions. Every one of the
644 AT rows is absent from the historian, so AT is safe to treat as a family.
But 73 FEV, 44 VEV and 3 CEV standalone rows have their own historian objects
and their own points - treating those families as parts would delete 120 real
controllable assets.
"""
import collections
import csv
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
IO = HERE / 'RDC_Historian_IO_list_CP2.xlsx'
MASTER = (HERE.parent / 'qnl-bms-room-allocation'
          / 'Appendix_A_Asset_Register_QNL_BMS_rooms.xlsx')
OUT = HERE / 'rdc_unit_or_part.csv'
RDC_FIRST = 1443


LEVELS = ('GF', '1F', '2F', 'MF', 'L0', 'B1')
FAMILY = 'AT|EV|VEV|CEV|FEV|EAV|VAV|CAV|FCU|AHU|ERU|EF'
# Only these can host a part. A fan coil unit shares a number with a terminal by
# coincidence - FCU1004 and VAV1004 are different units, drawn as separate
# widgets on South GF-6 - so letting FCU host gave 31 rows two candidate hosts
# and no way to choose. Air handling units are out for the same reason.
HOSTS = ('EV', 'VEV', 'CEV', 'FEV', 'EAV', 'VAV', 'CAV')


def strip_level(tag):
    """the tag with its level segments removed

    The register writes `RDC_NB_2F_2F_AHU8513` and the historian writes
    `RDC_NB_AHU8513` - the same unit under a different level segment, and one of
    them writes it twice. Comparing the whole string reported 26 air handling
    units as unknown to the historian when every one of them is in it.
    """
    out = re.sub(r'^(RDC_(?:NB_SB|NB|SB|UB))_', r'\1|', (tag or '').upper())
    while True:
        m = re.match(r'([^|]*\|)(?:%s)_(.+)$' % '|'.join(LEVELS), out)
        if not m:
            return out.replace('|', '_')
        out = m.group(1) + m.group(2)


def key(tag):
    return re.sub(r'[-_]', '', strip_level(tag))


def bld(tag):
    m = re.match(r'RDC_(NB_SB|NB|SB|UB)_', (tag or '').upper())
    return m.group(1) if m else '?'


def numbers(tag):
    """(family, number) for every unit named in a tag, level segments ignored"""
    body = re.sub(r'^RDC_(NB_SB|NB|SB|UB)_', '', strip_level(tag))
    return [(f, n.upper())
            for f, n in re.findall(r'(%s)-?(\d+[A-Za-z]?)' % FAMILY, body)]


def historian():
    """object -> how many points it carries"""
    wb = openpyxl.load_workbook(IO, read_only=True)
    count = collections.Counter()
    for name in wb.sheetnames:
        for r in wb[name].iter_rows(min_row=2, max_col=1, values_only=True):
            if r[0]:
                count[str(r[0]).split('.')[0]] += 1
    return count


def main():
    raw = historian()
    # keyed for looking a register tag up, and kept raw for reading the numbers
    # out of: `numbers` has to see the separators, or `VAV4110_EV4111` reads as
    # one unit numbered 4110E and the terminal AT-4110 finds no host.
    points = collections.Counter()
    for obj, n in raw.items():
        points[key(obj)] += n
    ws = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)[
        'Controllable Asset Registry']
    rows = [(i, str(r[0]).strip(), str(r[1] or '').strip(), str(r[3] or '').strip())
            for i, r in enumerate(ws.iter_rows(min_row=1, max_row=3202,
                                               max_col=4, values_only=True), 1)
            if r[0] and i >= RDC_FIRST]

    # every number the historian knows, and the object that carries it
    owner = {}
    for obj in raw:
        for fam, num in numbers(obj):
            if fam in HOSTS:
                # keyed on the building too: RDC_NB_1F_FCU2004 was matched to
                # RDC_SB_GF_VAV2004, a unit in the other building
                owner.setdefault((bld(obj), num), []).append(obj)

    out, tally = [], collections.Counter()
    for i, tag, kind, room in rows:
        k = key(tag)
        if points.get(k):
            verdict, parent = 'unit - has its own historian object', ''
        else:
            hosts = sorted({o for _, n in numbers(tag)
                            for o in owner.get((bld(tag), n), [])})
            hosts = [h for h in hosts if key(h) != k]
            if len(hosts) == 1:
                verdict = 'part - the historian carries the whole assembly'
                parent = hosts[0]
            elif hosts:
                verdict = 'part - but more than one host, needs a decision'
                parent = ' / '.join(hosts)
            else:
                verdict = 'neither - nothing in the historian knows this number'
                parent = ''
        tally[(kind, verdict)] += 1
        out.append([i, tag, kind, room, verdict, parent, points.get(k, 0)])

    with OUT.open('w', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['row', 'tag', 'type', 'room in column D', 'verdict',
                    'historian object it belongs to', 'points of its own'])
        w.writerows(out)

    wide = collections.Counter()
    for (kind, verdict), n in tally.items():
        wide[verdict] += n
    print('%-52s %s' % ('verdict', 'rows'))
    for v, n in wide.most_common():
        print('  %-50s %4d' % (v, n))
    print()
    print('by equipment type:')
    kinds = sorted({k for k, _ in tally})
    for kind in kinds:
        line = {v: n for (k, v), n in tally.items() if k == kind}
        print('  %-5s unit %4d | part %4d | unknown %3d'
              % (kind,
                 sum(n for v, n in line.items() if v.startswith('unit')),
                 sum(n for v, n in line.items() if v.startswith('part')),
                 sum(n for v, n in line.items() if v.startswith('neither'))))
    print('\nwrote %s' % OUT.name)


if __name__ == '__main__':
    main()
