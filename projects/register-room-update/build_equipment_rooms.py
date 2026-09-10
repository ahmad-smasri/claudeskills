#!/usr/bin/env python3
"""Equipment and its final room, one tab per building. Nothing else.

    python3 build_equipment_rooms.py

Three columns - tag, equipment type, room - on four tabs: HQ, QNL, SSC, RDC.
Every other column of the register is provenance: which source won, what the
BMS screen said, which drawing it came from. That belongs in the working file,
not in the sheet somebody reads to find out where a unit is.

Row order is the register's own, not alphabetical: the register groups by
equipment family and level, and sorting by tag would scatter every family.

A unit with no agreed room keeps an empty cell, shaded, rather than a guess.
"""
import argparse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SRC = HERE / 'Appendix_A_Asset_Register_clean.xlsx'
OUT = HERE / 'Appendix_A_Equipment_Rooms.xlsx'

BUILDINGS = ('HQ', 'QNL', 'SSC', 'RDC')
HEADERS = (('Tag No.', 30), ('Equipment type', 18), ('Room', 52))

TITLE_FILL = PatternFill('solid', fgColor='FF1F4E79')
HEAD_FILL = PatternFill('solid', fgColor='FF2E75B6')
NOROOM_FILL = PatternFill('solid', fgColor='FFFFF2CC')
WHITE = Font(bold=True, color='FFFFFFFF')
THIN = Side(style='thin', color='FFBFBFBF')
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical='top', wrap_text=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    ws = openpyxl.load_workbook(SRC, data_only=True)['Asset Register']
    rows = {b: [] for b in BUILDINGS}
    for r in range(3, ws.max_row + 1):
        bld = ws.cell(r, 12).value
        if bld not in rows:
            continue
        rows[bld].append((ws.cell(r, 1).value, ws.cell(r, 2).value,
                          ws.cell(r, 4).value))

    for b in BUILDINGS:
        blank = sum(1 for _, _, room in rows[b] if not room)
        print('%-4s %5d assets, %d with no agreed room' % (b, len(rows[b]), blank))
    print('     %5d total' % sum(len(v) for v in rows.values()))
    if args.dry_run:
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for b in BUILDINGS:
        sheet(wb.create_sheet(b), b, rows[b])
    wb.save(OUT)
    print('wrote %s' % OUT.name)


def sheet(ws, building, data):
    ws.cell(1, 1, '%s - EQUIPMENT AND ROOM' % building).font = Font(
        bold=True, size=13, color='FFFFFFFF')
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    for c in range(1, len(HEADERS) + 1):
        ws.cell(1, c).fill = TITLE_FILL
    for c, (name, width) in enumerate(HEADERS, 1):
        cell = ws.cell(2, c, name)
        cell.fill, cell.font, cell.border = HEAD_FILL, WHITE, GRID
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20

    for n, (tag, kind, room) in enumerate(data, 3):
        for c, value in enumerate((tag, kind, room), 1):
            cell = ws.cell(n, c, value)
            cell.alignment, cell.border = TOP, GRID
        if not room:
            ws.cell(n, 3).fill = NOROOM_FILL
    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = 'A2:C%d' % (len(data) + 2)


if __name__ == '__main__':
    main()
